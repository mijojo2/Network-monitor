import csv
import os
from typing import List, Tuple
from core.models import Device, SubDevice, get_default_sub_devices

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class ExcelService:
    """Service to export and import network devices to/from Excel (.xlsx) and CSV (.csv)."""

    @staticmethod
    def export_to_excel(devices: List[Device], file_path: str) -> Tuple[bool, str]:
        if not OPENPYXL_AVAILABLE:
            # Fallback to CSV if openpyxl is not installed
            csv_path = file_path if file_path.endswith(".csv") else file_path.rsplit(".", 1)[0] + ".csv"
            return ExcelService.export_to_csv(devices, csv_path)

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Network Devices"

            # Styles
            header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            online_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
            offline_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
            bold_font = Font(name="Segoe UI", size=10, bold=True)
            regular_font = Font(name="Segoe UI", size=10)

            headers = [
                "Device / Branch Name", "Parent IP", "Parent Status", "Latency",
                "Sub-Device Name", "Sub-Device IP", "Sub-Device Status", "Sub Latency"
            ]
            ws.append(headers)

            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            row_idx = 2
            for dev in devices:
                sub_count = len(dev.sub_devices)
                if sub_count == 0:
                    ws.append([dev.name, dev.ip, dev.status, dev.latency, "-", "-", "-", "-"])
                    row_idx += 1
                else:
                    for i, sub in enumerate(dev.sub_devices):
                        if i == 0:
                            ws.append([
                                dev.name, dev.ip, dev.status, dev.latency,
                                sub.name, sub.ip, sub.status, sub.latency
                            ])
                        else:
                            ws.append([
                                "", "", "", "",
                                sub.name, sub.ip, sub.status, sub.latency
                            ])

                        # Colorize status cells
                        p_stat_cell = ws.cell(row=row_idx, column=3)
                        if dev.status == "Online":
                            p_stat_cell.fill = online_fill
                        elif dev.status == "Offline":
                            p_stat_cell.fill = offline_fill

                        s_stat_cell = ws.cell(row=row_idx, column=7)
                        if sub.status == "Online":
                            s_stat_cell.fill = online_fill
                        elif sub.status == "Offline":
                            s_stat_cell.fill = offline_fill

                        row_idx += 1

            # Auto-fit column widths
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

            wb.save(file_path)
            return True, f"Successfully exported {len(devices)} devices to {os.path.basename(file_path)}"

        except Exception as e:
            return False, f"Export failed: {str(e)}"

    @staticmethod
    def export_to_csv(devices: List[Device], file_path: str) -> Tuple[bool, str]:
        try:
            with open(file_path, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f)
                writer.writerow([
                    "Branch Name", "Parent IP", "Parent Status", "Latency",
                    "Sub-Device Name", "Sub-Device IP", "Sub-Device Status", "Sub Latency"
                ])
                for dev in devices:
                    if not dev.sub_devices:
                        writer.writerow([dev.name, dev.ip, dev.status, dev.latency, "", "", "", ""])
                    else:
                        for i, sub in enumerate(dev.sub_devices):
                            if i == 0:
                                writer.writerow([
                                    dev.name, dev.ip, dev.status, dev.latency,
                                    sub.name, sub.ip, sub.status, sub.latency
                                ])
                            else:
                                writer.writerow([
                                    "", "", "", "",
                                    sub.name, sub.ip, sub.status, sub.latency
                                ])
            return True, f"Successfully exported to {os.path.basename(file_path)}"
        except Exception as e:
            return False, f"CSV Export failed: {str(e)}"

    @staticmethod
    def import_from_file(file_path: str) -> Tuple[bool, List[Device], str]:
        """Imports devices from an Excel (.xlsx) or CSV (.csv) file."""
        if not os.path.exists(file_path):
            return False, [], "File not found."

        if file_path.lower().endswith(".xlsx") and OPENPYXL_AVAILABLE:
            return ExcelService._import_xlsx(file_path)
        else:
            return ExcelService._import_csv(file_path)

    @staticmethod
    def _import_csv(file_path: str) -> Tuple[bool, List[Device], str]:
        try:
            devices_map = {}
            with open(file_path, "r", encoding="utf-8-sig") as f:
                reader = csv.reader(f)
                header = next(reader, None)
                current_device = None

                for row in reader:
                    if not row or not any(row):
                        continue
                    p_name = row[0].strip() if len(row) > 0 else ""
                    p_ip = row[1].strip() if len(row) > 1 else ""

                    if p_name and p_ip:
                        current_device = Device(name=p_name, ip=p_ip, sub_devices=[])
                        devices_map[p_ip] = current_device

                    # Check for sub-device in columns 4 & 5
                    sub_name = row[4].strip() if len(row) > 4 else ""
                    sub_ip = row[5].strip() if len(row) > 5 else ""

                    if current_device and sub_name and sub_ip:
                        current_device.sub_devices.append(SubDevice(name=sub_name, ip=sub_ip))

            devices = list(devices_map.values())
            for d in devices:
                if not d.sub_devices:
                    d.sub_devices = get_default_sub_devices(d.ip)

            return True, devices, f"Imported {len(devices)} devices successfully."
        except Exception as e:
            return False, [], f"CSV Import error: {str(e)}"

    @staticmethod
    def _import_xlsx(file_path: str) -> Tuple[bool, List[Device], str]:
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            ws = wb.active
            devices_map = {}
            current_device = None

            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return False, [], "Excel sheet is empty."

            for row in rows[1:]:  # skip header
                if not row or not any(row):
                    continue
                p_name = str(row[0]).strip() if len(row) > 0 and row[0] is not None else ""
                p_ip = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""

                if p_name and p_ip:
                    current_device = Device(name=p_name, ip=p_ip, sub_devices=[])
                    devices_map[p_ip] = current_device

                sub_name = str(row[4]).strip() if len(row) > 4 and row[4] is not None else ""
                sub_ip = str(row[5]).strip() if len(row) > 5 and row[5] is not None else ""

                if current_device and sub_name and sub_ip and sub_name != "-":
                    current_device.sub_devices.append(SubDevice(name=sub_name, ip=sub_ip))

            devices = list(devices_map.values())
            for d in devices:
                if not d.sub_devices:
                    d.sub_devices = get_default_sub_devices(d.ip)

            return True, devices, f"Imported {len(devices)} devices from Excel."
        except Exception as e:
            return False, [], f"Excel Import error: {str(e)}"
